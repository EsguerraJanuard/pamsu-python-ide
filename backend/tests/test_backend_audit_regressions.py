"""
Regression tests for the backend audit (2026-10-17).

Each test pins a defect that existed before the audit: password-reset
enumeration and synchronous delivery, unauthenticated or unbounded static
analysis, Judge0 callback status mapping, maintenance mode for existing
student tokens, admin input validation, dropped task fields, scheduled
publication, and two broken instructor endpoints.
"""

from datetime import datetime, timedelta, timezone

import bcrypt
import pytest
from sqlalchemy.orm import Session

from app.core import rate_limit
from app.core.security import (BCRYPT_ROUNDS, create_access_token,
                               get_password_hash, verify_password)
from app.core.system_settings import set_system_setting
from app.integrations import otp_delivery
from app.main import app
from app.models.domain_models import (AuditRecord, Classroom, Enrollment,
                                      ExecutionRequest, OTPChallenge, Task,
                                      User)

DOMAIN = "@pampangastateu.edu.ph"


def make_user(db: Session, *, email: str, role: str, school_id: str, name: str = "Test User", password: str = "Original#2026") -> User:
    user = User(
        name=name,
        email=email,
        role=role,
        school_id=school_id,
        password_hash=get_password_hash(password),
        email_verified=True,
        is_active=True,
    )
    db.add(user)
    db.commit()
    db.refresh(user)
    return user


def auth(user: User) -> dict[str, str]:
    return {"Authorization": f"Bearer {create_access_token({'sub': str(user.user_id)})}"}


@pytest.fixture
def people(db_session):
    admin = make_user(db_session, email=f"mis{DOMAIN}", role="admin", school_id="9000000001", name="MIS Admin")
    instructor = make_user(db_session, email=f"faculty{DOMAIN}", role="instructor", school_id="9000000002", name="Faculty One")
    other_instructor = make_user(db_session, email=f"faculty2{DOMAIN}", role="instructor", school_id="9000000003")
    student = make_user(db_session, email=f"2023000001{DOMAIN}", role="student", school_id="2023000001", name="Juan Dela Cruz")
    classroom = Classroom(
        instructor_id=instructor.user_id,
        name="Programming 1",
        subject_code="CCS101",
        section="BSIT 1A",
        class_code="CLASSA01",
        is_active=True,
    )
    db_session.add(classroom)
    db_session.commit()
    db_session.add(Enrollment(class_id=classroom.class_id, student_id=student.user_id, status="active"))
    db_session.commit()
    return {
        "admin": admin,
        "instructor": instructor,
        "other_instructor": other_instructor,
        "student": student,
        "classroom": classroom,
    }


# ---------------------------------------------------------------- password reset


class RecordingAdapter:
    def __init__(self):
        self.sent: list[tuple[str, str]] = []

    def send_otp(self, *, recipient_email, otp_code, purpose, expires_in_seconds):
        self.sent.append((recipient_email, otp_code))


@pytest.fixture
def otp_adapter():
    adapter = RecordingAdapter()
    app.dependency_overrides[otp_delivery.get_otp_delivery_adapter] = lambda: adapter
    yield adapter
    app.dependency_overrides.pop(otp_delivery.get_otp_delivery_adapter, None)


def start_reset(client, email: str):
    return client.post("/users/password-reset/start", json={"email": email})


def test_reset_start_response_is_identical_for_unknown_email(client, people, otp_adapter):
    known = start_reset(client, people["student"].email)
    unknown = start_reset(client, f"nobody{DOMAIN}")

    assert known.status_code == unknown.status_code == 201
    known_body, unknown_body = known.json(), unknown.json()
    assert set(known_body) == set(unknown_body)
    for field in ("purpose", "expires_in_seconds", "resend_after_seconds", "message"):
        assert known_body[field] == unknown_body[field]

    # Only the real account receives a code, and no response contains it.
    assert [email for email, _ in otp_adapter.sent] == [people["student"].email]
    code = otp_adapter.sent[0][1]
    assert code not in known.text


def test_reset_resend_for_unknown_email_behaves_like_a_real_challenge(client, db_session, people, otp_adapter):
    unknown = start_reset(client, f"nobody{DOMAIN}").json()

    challenge = db_session.query(OTPChallenge).filter(OTPChallenge.challenge_id == unknown["challenge_id"]).one()
    challenge.last_sent_at = datetime.now(timezone.utc) - timedelta(minutes=5)
    db_session.commit()

    response = client.post("/users/password-reset/resend", json={"challenge_id": unknown["challenge_id"]})

    assert response.status_code == 200
    assert otp_adapter.sent == []


def test_reset_completes_case_insensitively_and_revokes_old_tokens(client, db_session, people, otp_adapter):
    student = people["student"]
    old_headers = auth(student)
    challenge_id = start_reset(client, student.email.upper()).json()["challenge_id"]
    code = otp_adapter.sent[0][1]

    response = client.post(
        "/users/password-reset/complete",
        json={"challenge_id": challenge_id, "otp_code": code, "new_password": "BrandNew#2026"},
    )

    assert response.status_code == 200
    db_session.expire_all()
    assert verify_password("BrandNew#2026", db_session.get(User, student.user_id).password_hash)
    assert client.get("/activities/", headers=old_headers).status_code == 401
    # Single use.
    again = client.post(
        "/users/password-reset/complete",
        json={"challenge_id": challenge_id, "otp_code": code, "new_password": "Another#2026"},
    )
    assert again.status_code == 409


def test_new_reset_request_expires_the_previous_code(client, people, otp_adapter):
    first_id = start_reset(client, people["student"].email).json()["challenge_id"]
    first_code = otp_adapter.sent[0][1]
    start_reset(client, people["student"].email)

    response = client.post(
        "/users/password-reset/complete",
        json={"challenge_id": first_id, "otp_code": first_code, "new_password": "BrandNew#2026"},
    )

    assert response.status_code == 410


def test_reset_start_rejects_non_university_email(client, otp_adapter):
    assert start_reset(client, "someone@gmail.com").status_code == 422


# ---------------------------------------------------------------- rate limiting


class FakeRedis:
    def __init__(self):
        self.values: dict[str, int] = {}

    def incr(self, key):
        self.values[key] = self.values.get(key, 0) + 1
        return self.values[key]

    def expire(self, key, seconds):
        return True

    def ttl(self, key):
        return 30


def test_rate_limit_is_per_identifier_and_fails_open(monkeypatch):
    fake = FakeRedis()
    monkeypatch.setattr(rate_limit, "redis_client", fake)

    for _ in range(3):
        rate_limit.hit_rate_limit(scope="t", identifier="user-1", limit=3, window_seconds=60)

    with pytest.raises(rate_limit.RateLimitExceededError):
        rate_limit.hit_rate_limit(scope="t", identifier="user-1", limit=3, window_seconds=60)

    rate_limit.hit_rate_limit(scope="t", identifier="user-2", limit=3, window_seconds=60)
    assert all("user-1" not in key for key in fake.values)

    class DownRedis:
        def incr(self, key):
            raise ConnectionError("down")

    monkeypatch.setattr(rate_limit, "redis_client", DownRedis())
    rate_limit.hit_rate_limit(scope="t", identifier="user-1", limit=0, window_seconds=60)


def test_password_reset_start_is_limited_per_email(client, people, otp_adapter, monkeypatch):
    monkeypatch.setattr(rate_limit, "redis_client", FakeRedis())

    statuses = [start_reset(client, f"nobody{DOMAIN}").status_code for _ in range(4)]

    assert statuses == [201, 201, 201, 429]


def test_new_password_hashes_use_capped_bcrypt_rounds():
    assert BCRYPT_ROUNDS <= 8
    assert bcrypt.gensalt(rounds=BCRYPT_ROUNDS)[4:6] == get_password_hash("whatever1")[4:6].encode()


# ---------------------------------------------------------------- static analysis


def test_lint_requires_authentication(client, people):
    assert client.post("/execution/lint", json={"code": "x = 1\n"}).status_code == 401
    response = client.post("/execution/lint", headers=auth(people["student"]), json={"code": "import os\n"})
    assert response.status_code == 200
    assert response.json()["markers"][0]["severity"] == "warning"


# Short ids: pytest exports the test id in PYTEST_CURRENT_TEST, and Windows
# rejects environment variables longer than 32767 characters.
@pytest.mark.parametrize(
    "code",
    ["-" * 99_999 + "1", "a+" * 40_000 + "a"],
    ids=["unary-minus-chain", "long-binop-chain"],
)
def test_pathological_code_is_reported_not_crashing(client, people, code):
    response = client.post("/execution/lint", headers=auth(people["student"]), json={"code": code})

    assert response.status_code == 200
    assert "too" in response.json()["markers"][0]["message"]

    analysis = client.post(
        "/instructors/tasks/analyze-solution",
        headers=auth(people["instructor"]),
        json={"reference_code": code},
    )
    assert analysis.status_code == 400


def test_static_analysis_input_is_bounded(client, people):
    huge = "x = 1\n" * 20_000

    assert client.post("/execution/lint", headers=auth(people["student"]), json={"code": huge}).status_code == 422
    response = client.post(
        "/instructors/tasks/analyze-solution",
        headers=auth(people["instructor"]),
        json={"reference_code": huge},
    )
    assert response.status_code == 422


# ---------------------------------------------------------------- Judge0 callback


def make_execution_request(db: Session, student: User) -> ExecutionRequest:
    request = ExecutionRequest(
        execution_id="11111111-1111-4111-8111-111111111111",
        correlation_id="22222222-2222-4222-8222-222222222222",
        student_id=student.user_id,
        request_kind="run",
        status="running",
        source_code="print(1)\n",
        standard_input="",
    )
    db.add(request)
    db.commit()
    return request


@pytest.mark.parametrize(
    ("status_payload", "expected"),
    [
        ({"id": 3}, "completed"),
        ({"id": 5}, "timed_out"),
        ({"id": 11}, "runtime_error"),
        ({"id": 2}, "failed"),
        ({"id": "3"}, "failed"),
        (None, "failed"),
    ],
)
def test_judge0_callback_status_mapping(client, db_session, people, status_payload, expected):
    record = make_execution_request(db_session, people["student"])

    response = client.put(
        "/execution/internal/judge0-callback",
        params={"execution_id": record.execution_id, "correlation_id": record.correlation_id},
        json={"status": status_payload, "stdout": None},
    )

    assert response.status_code == 200, response.text
    db_session.expire_all()
    assert db_session.get(ExecutionRequest, record.execution_id).status == expected


def test_judge0_callback_requires_matching_correlation_id(client, db_session, people):
    record = make_execution_request(db_session, people["student"])

    response = client.put(
        "/execution/internal/judge0-callback",
        params={"execution_id": record.execution_id, "correlation_id": "33333333-3333-4333-8333-333333333333"},
        json={"status": {"id": 3}},
    )

    assert response.status_code == 404
    db_session.expire_all()
    assert db_session.get(ExecutionRequest, record.execution_id).status == "running"


# ---------------------------------------------------------------- maintenance mode


def test_maintenance_mode_blocks_existing_student_tokens_only(client, db_session, people):
    student_headers = auth(people["student"])
    assert client.get("/activities/", headers=student_headers).status_code == 200

    set_system_setting(db_session, "maintenance_mode", True)

    assert client.get("/activities/", headers=student_headers).status_code == 503
    assert client.get("/classrooms/", headers=auth(people["instructor"])).status_code == 200


# ---------------------------------------------------------------- admin endpoints


@pytest.mark.parametrize(
    "payload",
    [
        {"email": "x@gmail.com", "first_name": "A", "last_name": "B", "school_id": "2023000002", "password": "Temporary#1"},
        {"email": f"a{DOMAIN}", "first_name": "A", "last_name": "B", "school_id": "12", "password": "Temporary#1"},
        {"email": f"a{DOMAIN}", "first_name": "A", "last_name": "B", "school_id": "2023000002", "password": "short"},
        {"email": f"a{DOMAIN}", "first_name": "", "last_name": "B", "school_id": "2023000002", "password": "Temporary#1"},
    ],
)
def test_admin_student_creation_validates_input(client, people, payload):
    assert client.post("/admin/students", headers=auth(people["admin"]), json=payload).status_code == 422


def test_admin_student_email_is_normalized_and_unique_case_insensitively(client, db_session, people):
    payload = {"email": f"New.Student{DOMAIN}", "first_name": "New", "last_name": "Student", "school_id": "2023000002", "password": "Temporary#1"}

    created = client.post("/admin/students", headers=auth(people["admin"]), json=payload)
    duplicate = client.post(
        "/admin/students",
        headers=auth(people["admin"]),
        json={**payload, "email": f"NEW.STUDENT{DOMAIN}", "school_id": "2023000003"},
    )

    assert created.status_code == 200
    assert created.json()["email"] == f"new.student{DOMAIN}"
    assert duplicate.status_code == 400


def test_admin_cannot_hard_delete_account_with_records(client, db_session, people):
    response = client.delete(f"/admin/users/{people['instructor'].user_id}", headers=auth(people["admin"]))

    assert response.status_code == 409
    db_session.expire_all()
    assert db_session.get(User, people["instructor"].user_id) is not None


def test_admin_password_reset_and_list_inputs_are_bounded(client, people):
    headers = auth(people["admin"])

    assert client.patch(f"/admin/users/{people['student'].user_id}/password", headers=headers, json={"new_password": "x"}).status_code == 422
    assert client.get("/admin/users", headers=headers, params={"limit": 100_000}).status_code == 422
    assert client.get("/admin/audit-logs", headers=headers, params={"skip": -1}).status_code == 422


def test_admin_settings_reject_unknown_strictness(client, people):
    response = client.patch(
        "/admin/settings",
        headers=auth(people["admin"]),
        json={"maintenance_mode": False, "default_ast_strictness": "anything", "registration_enabled": False},
    )

    assert response.status_code == 422


def test_profile_rejects_unknown_strictness(client, people):
    response = client.patch(
        "/users/me",
        headers=auth(people["instructor"]),
        json={"name": "Faculty One", "ast_strictness_level": "extreme"},
    )

    assert response.status_code == 422


# ---------------------------------------------------------------- tasks


def create_task(client, people, **extra):
    response = client.post(
        "/instructors/tasks/",
        headers=auth(people["instructor"]),
        json={"title": "Loops", "class_ids": [people["classroom"].class_id], **extra},
    )
    assert response.status_code == 201, response.text
    body = response.json()
    return body[0] if isinstance(body, list) else body


def test_task_creation_keeps_difficulty_and_schedule(client, db_session, people):
    scheduled = (datetime.now(timezone.utc) + timedelta(days=1)).isoformat()

    body = create_task(client, people, difficulty="expert", scheduled_publish_at=scheduled)

    task = db_session.get(Task, body["task_id"])
    assert task.difficulty == "expert"
    assert task.scheduled_publish_at is not None
    assert task.is_published is False


def test_scheduled_task_publishes_when_due_and_stays_unpublished_after_unpublish(client, db_session, people):
    body = create_task(
        client,
        people,
        scheduled_publish_at=(datetime.now(timezone.utc) + timedelta(days=1)).isoformat(),
    )
    task = db_session.get(Task, body["task_id"])
    student_headers = auth(people["student"])

    assert client.get("/activities/", headers=student_headers).json() == []

    task.scheduled_publish_at = datetime.now(timezone.utc) - timedelta(minutes=1)
    db_session.commit()

    listed = client.get("/activities/", headers=student_headers).json()
    assert [item["task_id"] for item in listed] == [task.task_id]
    db_session.expire_all()
    assert db_session.query(AuditRecord).filter(
        AuditRecord.action_type == "activity_published",
        AuditRecord.actor_user_id == people["instructor"].user_id,
    ).count() == 1

    unpublish = client.patch(
        f"/instructors/tasks/{task.task_id}/publication",
        headers=auth(people["instructor"]),
        json={"is_published": False},
    )
    assert unpublish.status_code == 200, unpublish.text
    assert client.get("/activities/", headers=student_headers).json() == []


# ---------------------------------------------------------------- instructor endpoints


def test_instructor_growth_analytics_for_enrolled_student(client, people):
    path = f"/instructors/students/{people['student'].user_id}/analytics/growth"

    assert client.get(path, headers=auth(people["instructor"])).status_code == 200
    assert client.get(path, headers=auth(people["other_instructor"])).status_code == 403


def test_review_queue_export_is_owner_scoped_csv(client, people):
    response = client.get("/instructors/review-queue/export", headers=auth(people["instructor"]))

    assert response.status_code == 200
    assert response.text.splitlines()[0].startswith("Submission ID,Student Name,School ID")

    foreign = client.get(
        "/instructors/review-queue/export",
        headers=auth(people["other_instructor"]),
        params={"class_id": people["classroom"].class_id},
    )
    assert foreign.status_code == 403


def test_guest_can_update_own_profile(client):
    guest = client.post("/guest")
    assert guest.status_code == 200, guest.text
    headers = {"Authorization": f"Bearer {guest.json()['access_token']}"}

    response = client.patch("/users/me", headers=headers, json={"name": "Visiting Student"})

    assert response.status_code == 200, response.text
    assert response.json()["school_id"].startswith("GST")


def test_gradebook_xlsx_does_not_turn_student_names_into_formulas(client, db_session, people):
    import io

    from openpyxl import load_workbook

    from app.models.domain_models import InstructorGrade, Submission

    student = people["student"]
    student.name = '=HYPERLINK("https://example.invalid","Open")'
    task = Task(
        class_id=people["classroom"].class_id,
        instructor_id=people["instructor"].user_id,
        title="Loops",
        starter_code="",
        is_published=True,
    )
    db_session.add(task)
    db_session.commit()
    submission = Submission(
        student_id=student.user_id,
        task_id=task.task_id,
        attempt_number=1,
        raw_code="print(1)\n",
        standard_input="",
        status="graded",
        is_official=True,
    )
    db_session.add(submission)
    db_session.commit()
    db_session.add(
        InstructorGrade(
            submission_id=submission.sub_id,
            instructor_id=people["instructor"].user_id,
            score=90,
            max_score=100,
            is_released=False,
        )
    )
    db_session.commit()

    response = client.get(
        f"/reports/classrooms/{people['classroom'].class_id}/gradebook.xlsx",
        headers=auth(people["instructor"]),
    )

    assert response.status_code == 200, response.text
    sheet = load_workbook(io.BytesIO(response.content)).active
    names = [row[0].value for row in sheet.iter_rows(min_row=2)]
    assert names == ["'" + student.name]
    assert all(row[0].data_type != "f" for row in sheet.iter_rows(min_row=2))
